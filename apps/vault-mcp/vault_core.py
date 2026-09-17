# vault_core.py -- the actual vault/OCR/chunking/Qdrant logic, shared by
# server.py (the MCP server Hermes talks to) and upload_web.py (the new
# plain web-upload page, 2026-09-17). Pulled out of server.py unchanged --
# same functions, same behavior, just importable from two processes instead
# of living only in the MCP one. Nothing in here talks MCP at all; it's
# plain Python + httpx + the file-extraction libraries.
#
# Why two processes instead of one: mcp's MCPServer.run() owns its own
# event loop/ASGI app and there's no documented way to mount extra plain
# HTTP routes onto it, so bolting a browser-facing upload form directly
# onto server.py would mean fighting an undocumented internal API on a
# server Hermes depends on in production. Splitting into two small
# processes that both import this module is safer: server.py (MCP, port
# 8000) is untouched behavior-wise, and upload_web.py (plain FastAPI, port
# 8081) is a completely separate, low-stakes thing to get wrong. See
# entrypoint.sh for how both get started in the one container.
import base64
import hashlib
import logging
import os
import re
import shutil
import uuid
from pathlib import Path

import httpx

logging.basicConfig(level=logging.INFO, format="%(asctime)s vault-core %(levelname)s: %(message)s")
logger = logging.getLogger("vault-core")

QDRANT_URL = os.environ.get("QDRANT_URL", "http://qdrant.ai-memory.svc.cluster.local:6333")
COLLECTION_NAME = "obsidian_vault"

LITELLM_URL = os.environ.get("LITELLM_URL", "http://litellm.ai-models.svc.cluster.local:4000/v1")
LITELLM_API_KEY = os.environ.get("LITELLM_API_KEY", "")
EMBED_MODEL = "nomic-embed-text-v1.5"
VISION_MODEL = "qwen3.8-27b-Expert-Vision"   # same locked vision route Hermes's own image-in-chat path already uses (custom_router_hook.py's _has_image check)

VAULT_PATH = Path(os.environ.get("VAULT_PATH", "/data/obsidian_vault"))
INBOX_DIR = VAULT_PATH / "Inbox"             # default folder when write_note/ingest_attachment's folder param is left unset
ATTACHMENTS_DIR = VAULT_PATH / "Attachments"

# Folders write_note/ingest_attachment's folder param is allowed to route a
# note into -- a fixed whitelist, not any string a caller sends, so a
# hallucinated/typo'd folder name (from Hermes) or a bad form value (from
# the web upload page) can't scatter notes outside Eldar's real top-level
# vault structure or silently create a new folder he never asked for.
# Matches Eldar's actual Obsidian top-level folders as of 2026-09-17 -- if
# he adds a new one in Obsidian, add it here too.
#
# Expanded 2026-09-17 with Work, Finance, Transport, Travels, Education (was
# just Inbox/Family/Health/Legal/Household/Reference/AI). A few of these
# overlap in ways worth spelling out so routing stays consistent:
#   - Work vs Legal: employment docs, contracts, B2B/consulting -> Work.
#     Legal stays for residency, IDs, property, agreements not tied to a job.
#   - Finance vs Transport: ALL insurance (car, home, life, health) -> Finance,
#     even car insurance -- Transport is for vehicle-specific operational
#     docs (registration, purchase, tickets, service records), not insurance.
ALLOWED_FOLDERS = {
    "Inbox", "Family", "Health", "Legal", "Household", "Reference", "AI",
    "Work", "Finance", "Transport", "Travels", "Education",
}

# Both server.py and upload_web.py run as root in their container (see
# Dockerfile -- no USER set), same as most single-pod services on this
# cluster, but hermes-agent's own gateway process deliberately drops to a
# non-root "hermes" user (uid/gid 10000, see k8s/p3/hermes-agent.yaml) once
# it boots. That mismatch is real: found 2026-09-14 when Hermes could copy a
# photo into Attachments/ (its own writes, as itself) but then couldn't edit
# or delete the .md notes THIS server had written moments earlier, because
# they came out owned root:root and Hermes has no sudo in its container.
# Keeping this GID here as a plain constant (not read from the environment)
# is a deliberate shortcut, not an oversight -- see server.py's original
# comment (now here) for the full reasoning. If hermes-agent's uid ever
# changes, update it here too.
HERMES_GID = 10000


def chgrp_writable(path: Path, dir_mode=0o775, file_mode=0o664):
    # Group-owns path by Hermes's gid and makes it group-writable, so the
    # non-root hermes-agent process can edit or delete something a root
    # process (this one, or upload_web.py) just created. Root can chown to
    # any group regardless of who currently owns the file, so this is safe
    # to call unconditionally.
    mode = dir_mode if path.is_dir() else file_mode
    os.chown(path, -1, HERMES_GID)   # -1 leaves the user (root) unchanged, only the group changes
    path.chmod(mode)


# Created at import time, not lazily -- runs once per process (server.py
# AND upload_web.py each import this module, so this runs twice total, once
# per process start). Harmless and deliberately idempotent (mkdir
# exist_ok=True, chgrp unconditional-but-safe) -- also means either process
# restarting alone heals any folder that drifted back to root-only
# permissions, same reasoning as the original single-process version of
# this comment.
for _folder_name in ALLOWED_FOLDERS:
    _folder_path = VAULT_PATH / _folder_name
    _folder_path.mkdir(parents=True, exist_ok=True)
    chgrp_writable(_folder_path)
ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
chgrp_writable(ATTACHMENTS_DIR)

# Where ingest_file (server.py, MCP path) is allowed to read an existing
# path from -- see vault-mcp's original top-of-file comment for why it's
# just these two (both volumes the pod fully owns, no cross-pod reads into
# someone else's storage). Not used by upload_web.py's own upload path --
# a browser upload has no server-side path to validate, it hands over raw
# bytes instead (see save_upload_to_attachments below).
ALLOWED_INPUT_ROOTS = [
    VAULT_PATH,
    Path(os.environ.get("SHARED_INPUT_DIR", "/shared-output")),
]

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}

# Hard cap on how many pages of a scanned/no-text-layer PDF get OCR'd via the
# vision model. The vision model route is locked --parallel 1 (see
# /areas/ai-box-k3s.md), so pages OCR this way are strictly sequential --
# one page can take 30-90s. Past this cap the note still gets the pages that
# were OCR'd plus a note that it was truncated, rather than the whole call
# failing/timing out.
MAX_OCR_PDF_PAGES = 3

# ---- chunking -- identical to scripts/vault_sync.py's own copy of this
# (see vault-mcp's original top comment for why it's duplicated there, not
# imported -- that reasoning is about the host-process/container boundary
# and still applies; it has nothing to do with this module, which both
# vault-mcp's own processes share normally) ----
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
FENCE_RE = re.compile(r"^(```|~~~)")
TABLE_RE = re.compile(r"^\s*\|.*\|\s*$")
LIST_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
FRONTMATTER_RE = re.compile(r"^---\s*\n.*?\n---\s*\n", re.DOTALL)


def _flush_block(blocks, lines, block_type, headings):
    text = "\n".join(lines).strip()
    if text:
        blocks.append({"text": text, "chunk_type": block_type, "heading_path": " > ".join(headings)})
    lines.clear()


def blockify(markdown):
    blocks = []
    current_lines = []
    current_type = "paragraph"
    headings = []
    in_code = False

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if FENCE_RE.match(stripped):
            if in_code:
                current_lines.append(line)
                _flush_block(blocks, current_lines, "code", headings)
                current_type = "paragraph"
                in_code = False
            else:
                _flush_block(blocks, current_lines, current_type, headings)
                current_lines.append(line)
                current_type = "code"
                in_code = True
            continue

        if in_code:
            current_lines.append(line)
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            _flush_block(blocks, current_lines, current_type, headings)
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            headings = headings[: level - 1] + [title]
            blocks.append({"text": stripped, "chunk_type": "heading", "heading_path": " > ".join(headings)})
            current_type = "paragraph"
            continue

        if not stripped:
            _flush_block(blocks, current_lines, current_type, headings)
            current_type = "paragraph"
            continue

        line_type = "table" if TABLE_RE.match(line) else "list" if LIST_RE.match(line) else "paragraph"
        if current_lines and line_type != current_type:
            _flush_block(blocks, current_lines, current_type, headings)
        current_type = line_type
        current_lines.append(line)

    _flush_block(blocks, current_lines, current_type, headings)
    return blocks


def chunk_hash(text: str) -> str:
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def _slugify(title: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", title.strip().lower()).strip("-")
    return slug or "untitled"


# sends this to LiteLLM's OpenAI-compatible embeddings endpoint:
#   {"model": "nomic-embed-text-v1.5", "input": ["chunk one text", "chunk two text"]}
# returns a list of vectors in the same order as texts
# nomic-embed-text-v1.5 wants a "search_document: " prefix on text being
# indexed and "search_query: " on text a search is run with -- without it,
# vectors cluster too close together to tell related notes from unrelated
# ones (see scripts/vault_sync.py's embed_texts for the full story). prefix
# must be "search_document" or "search_query" -- callers pass one
# explicitly rather than defaulting, so it's never accidentally left off.
async def embed_texts(client: httpx.AsyncClient, texts: list[str], prefix: str) -> list[list[float]]:
    headers = {"Authorization": f"Bearer {LITELLM_API_KEY}"}
    resp = await client.post(
        f"{LITELLM_URL}/embeddings",
        json={"model": EMBED_MODEL, "input": [f"{prefix}: {t}" for t in texts]},
        headers=headers,
        timeout=60.0,
    )
    resp.raise_for_status()
    result = resp.json()
    ordered = sorted(result["data"], key=lambda d: d["index"])
    return [d["embedding"] for d in ordered]


# ---- Qdrant, plain REST -- no qdrant-client dependency, same reasoning as
# vault_sync.py dropping langchain: this container already pulls in httpx
# for LiteLLM, no reason to add a second HTTP-wrapping SDK just for Qdrant ----

async def qdrant_ensure_collection(client: httpx.AsyncClient, vector_size: int) -> None:
    resp = await client.get(f"{QDRANT_URL}/collections/{COLLECTION_NAME}")
    if resp.status_code == 200:
        return
    await client.put(
        f"{QDRANT_URL}/collections/{COLLECTION_NAME}",
        json={"vectors": {"size": vector_size, "distance": "Cosine"}},
        timeout=30.0,
    )
    logger.info("created Qdrant collection '%s' (size=%s)", COLLECTION_NAME, vector_size)


async def qdrant_delete_source(client: httpx.AsyncClient, source: str) -> None:
    await client.post(
        f"{QDRANT_URL}/collections/{COLLECTION_NAME}/points/delete",
        json={"filter": {"must": [{"key": "source", "match": {"value": source}}]}},
        timeout=30.0,
    )


async def qdrant_upsert(client: httpx.AsyncClient, points: list[dict]) -> None:
    resp = await client.put(
        f"{QDRANT_URL}/collections/{COLLECTION_NAME}/points",
        json={"points": points},
        timeout=60.0,
    )
    resp.raise_for_status()


# sends this to Qdrant's search endpoint:
#   {"vector": [0.01, ...], "limit": 5, "with_payload": true, "filter": {...}}
# returns e.g. {"result": [{"id": "...", "score": 0.82, "payload": {...}}, ...]}
async def qdrant_search(client: httpx.AsyncClient, vector: list[float], top_k: int, tags) -> list[dict]:
    body = {"vector": vector, "limit": top_k, "with_payload": True}
    if tags:
        body["filter"] = {"must": [{"key": "tags", "match": {"any": tags}}]}
    resp = await client.post(
        f"{QDRANT_URL}/collections/{COLLECTION_NAME}/points/search",
        json=body,
        timeout=30.0,
    )
    resp.raise_for_status()
    return resp.json()["result"]


# ---- writing + indexing a vault note -- shared by ingest_note and
# ingest_attachment, both of which just build the body text differently and
# hand it here ----

def write_note(title: str, tags, audience, body: str, folder: str = "Inbox") -> tuple[Path, str]:
    if folder not in ALLOWED_FOLDERS:
        raise ValueError(f"folder must be one of {sorted(ALLOWED_FOLDERS)} (got {folder!r})")
    target_dir = VAULT_PATH / folder
    target_dir.mkdir(parents=True, exist_ok=True)
    slug = _slugify(title)
    path = target_dir / f"{slug}.md"
    if path.exists():   # don't clobber an earlier note with the same title -- give this one its own file
        path = target_dir / f"{slug}-{uuid.uuid4().hex[:8]}.md"

    front_lines = []
    if tags:
        front_lines.append("tags: [" + ", ".join(tags) + "]")
    if audience:
        front_lines.append(f"audience: {audience}")
    frontmatter = "---\n" + "\n".join(front_lines) + "\n---\n\n" if front_lines else ""

    path.write_text(f"{frontmatter}# {title}\n\n{body}\n", encoding="utf-8")
    chgrp_writable(path)   # otherwise Hermes (non-root) can't edit or delete the note this (root) process just wrote
    source = str(path.relative_to(VAULT_PATH))
    return path, source


async def index_note(client: httpx.AsyncClient, path: Path, source: str, title: str, tags, audience) -> int:
    text = path.read_text(encoding="utf-8")
    body = FRONTMATTER_RE.sub("", text, count=1)   # frontmatter itself isn't searchable content
    blocks = blockify(body)
    if not blocks:
        return 0

    vectors = await embed_texts(client, [b["text"] for b in blocks], "search_document")

    points = []
    for block, vector in zip(blocks, vectors):
        h = chunk_hash(block["text"])
        points.append({
            # same source+hash keying as vault_sync.py -- see vault-mcp's
            # original top comment for the collision reasoning
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source}::{h}")),
            "vector": vector,
            "payload": {
                "text": block["text"],
                "chunk_type": block["chunk_type"],
                "heading_path": block["heading_path"],
                "source": source,
                "note_title": title,
                "tags": tags or [],
                "audience": audience,
                "chunk_hash": h,
            },
        })

    await qdrant_ensure_collection(client, len(vectors[0]))
    await qdrant_delete_source(client, source)
    await qdrant_upsert(client, points)
    return len(points)


# Last-resort OCR when the vision model call itself fails (LiteLLM down/
# unreachable) -- classic tesseract via pytesseract, no LLM involved. Weaker
# than the vision model on skewed photos or handwriting, but doesn't depend
# on anything else being up.
def _ocr_image_locally(image_bytes: bytes) -> str:
    import io
    import pytesseract
    from PIL import Image
    image = Image.open(io.BytesIO(image_bytes))
    return pytesseract.image_to_string(image).strip()


# sends this to LiteLLM's vision-capable model:
#   {"model": "qwen3.8-27b-Expert-Vision", "messages": [{"role": "user",
#    "content": [{"type": "text", "text": "..."}, {"type": "image_url",
#    "image_url": {"url": "data:image/png;base64,..."}}]}], "max_tokens": 2000}
# returns the usual OpenAI chat shape: {"choices": [{"message": {"content": "..."}}]}
#
# Prompted for verbatim transcription, not a scene caption -- the real use
# case here is photographed notes/documents (including things like
# prescriptions) Eldar wants findable by their actual content, not "a photo
# of a piece of paper". Falls back to local pytesseract (see above) if the
# vision model call itself fails.
async def ocr_image(client: httpx.AsyncClient, image_bytes: bytes, mime_type: str) -> str:
    b64 = base64.b64encode(image_bytes).decode("ascii")
    headers = {"Authorization": f"Bearer {LITELLM_API_KEY}"}
    prompt = (
        "Transcribe ALL visible text in this image, verbatim and in full, exactly as written "
        "(keep the original language -- don't translate it). Preserve line breaks and structure "
        "(headings, lists, tables) as plain text where it's clear. If there is little or no text, "
        "instead describe what the image shows in one or two sentences."
    )
    try:
        resp = await client.post(
            f"{LITELLM_URL}/chat/completions",
            headers=headers,
            json={
                "model": VISION_MODEL,
                "messages": [{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64}"}},
                    ],
                }],
                "max_tokens": 2000,
            },
            timeout=90.0,
        )
        resp.raise_for_status()
        return resp.json()["choices"][0]["message"]["content"].strip()
    except httpx.HTTPError as exc:
        logger.warning("vision OCR failed (%s) -- falling back to local pytesseract", exc)
        return _ocr_image_locally(image_bytes)


# FIXED 2026-09-17: was pypdf's PdfReader.extract_text(). Real bug found
# ingesting Eldar's PMI job certificate through the new web upload page --
# pypdf split the extracted text into ONE CHARACTER PER LINE ("d\no\ne\ns"
# instead of "does"), a known pypdf weakness on PDFs with certain font/glyph
# encodings (this one's Polish-issued, likely an uncommon embedded font
# setup). The text still made it into the vault note and got embedded into
# Qdrant, just shredded into meaningless single-letter chunks -- nothing
# crashed, nothing errored, it just wasn't searchable, which is why the
# note looked "not there" from Hermes's side even though it genuinely was.
# Confirmed directly via Qdrant's own points/scroll API, not just inferred.
# Switched to pymupdf (fitz) instead -- already a hard dependency here (used
# for the scanned-PDF OCR fallback below), so this isn't a new dependency,
# just using the more reliable extractor that was already in the container.
# pypdf itself is no longer used anywhere in this file; dropped from
# requirements.txt.
def extract_pdf_text(path: Path) -> str:
    import fitz   # pymupdf
    doc = fitz.open(str(path))
    try:
        return "\n\n".join(page.get_text() for page in doc).strip()
    finally:
        doc.close()


# Only meaningful on a PDF with a real text layer -- extract_pdf_text()
# above flattens a table into hard-to-read run-on text, so any tables camelot
# finds get appended separately as plain pipe rows. camelot's "lattice"
# flavor (the default) needs visible table borders/lines to find a table;
# if a real PDF turns out to have borderless tables, switch to
# flavor="stream" (whitespace-based instead).
def extract_pdf_tables(path: Path) -> str:
    import camelot
    tables = camelot.read_pdf(str(path), pages="all")
    if len(tables) == 0:
        return ""
    parts = []
    for i, table in enumerate(tables):
        rows = table.df.values.tolist()
        lines = [" | ".join(str(cell).strip() for cell in row) for row in rows]
        parts.append(f"Table {i + 1} (page {table.page}):\n" + "\n".join(lines))
    return "\n\n".join(parts)


# Fallback for a PDF with no extractable text layer (scanned/photographed).
# Rasterizes each page to a PNG via pymupdf (no external system binary
# needed, unlike poppler/pdf2image) and runs it through the same vision-OCR
# path images use. Capped at MAX_OCR_PDF_PAGES -- see that constant's
# comment for why (the vision route is strictly sequential).
async def extract_pdf_via_ocr(client: httpx.AsyncClient, path: Path) -> str:
    import fitz   # pymupdf
    doc = fitz.open(str(path))
    try:
        page_count = doc.page_count
        pages_to_ocr = min(page_count, MAX_OCR_PDF_PAGES)
        texts = []
        for i in range(pages_to_ocr):
            png_bytes = doc[i].get_pixmap(dpi=200).tobytes("png")
            texts.append(await ocr_image(client, png_bytes, "image/png"))
        result = "\n\n---\n\n".join(t for t in texts if t)
        if page_count > MAX_OCR_PDF_PAGES:
            result += f"\n\n[OCR truncated -- {page_count - MAX_OCR_PDF_PAGES} more page(s) not processed]"
        return result
    finally:
        doc.close()


# docx is already digital text -- no OCR needed, just walk paragraphs and
# tables in document order. Simple pipe-row rendering for tables (not a
# proper markdown table, but good enough to be searchable and readable).
def extract_docx_text(path: Path) -> str:
    import docx   # python-docx
    document = docx.Document(str(path))
    parts = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


# Picks a collision-safe destination filename under Attachments/ for a new
# file named `filename` -- shared by server.py's ingest_file (copying an
# existing path) and upload_web.py's upload endpoint (writing fresh bytes).
def _unique_attachment_path(filename: str) -> Path:
    dest = ATTACHMENTS_DIR / filename
    if dest.exists():
        stem, suffix = Path(filename).stem, Path(filename).suffix
        dest = ATTACHMENTS_DIR / f"{stem}-{uuid.uuid4().hex[:8]}{suffix}"
    return dest


# Used by server.py's ingest_file: src is a path Hermes already gave us
# (already under VAULT_PATH or SHARED_INPUT_DIR). Copies it into
# Attachments/ unless it's already there.
def copy_to_attachments(src: Path) -> Path:
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    if src.parent == ATTACHMENTS_DIR:
        return src   # already placed there -- don't duplicate it under a new name
    dest = _unique_attachment_path(src.name)
    shutil.copy2(src, dest)
    chgrp_writable(dest)   # copy2 preserves src's own mode/owner (root's, here) -- fix it up so Hermes can still touch its own attachment later
    return dest


# Used by upload_web.py: writes freshly-uploaded bytes straight into
# Attachments/ (there's no existing on-disk path to copy from -- the
# browser handed over the file's bytes directly).
def save_upload_to_attachments(filename: str, data: bytes) -> Path:
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    dest = _unique_attachment_path(filename)
    dest.write_bytes(data)
    chgrp_writable(dest)
    return dest


# The shared "extract text by file type, write the vault note, index it"
# logic -- this is the exact body of the original single-file ingest_file
# tool, minus the part that gets the file INTO Attachments/ (callers do
# that their own way first: server.py's ingest_file copies an existing
# path via copy_to_attachments, upload_web.py writes fresh bytes via
# save_upload_to_attachments) -- see this module's top comment for why.
#
# dest must already be the final file living under ATTACHMENTS_DIR.
# Returns a dict: {"status": ..., "source": "<folder>/<slug>.md",
# "attachment_ref": "Attachments/<name>", "chunk_count": int}
async def ingest_attachment(client: httpx.AsyncClient, dest: Path, title: str | None, tags, audience, folder: str) -> dict:
    import mimetypes
    if folder not in ALLOWED_FOLDERS:
        raise ValueError(f"folder must be one of {sorted(ALLOWED_FOLDERS)} (got {folder!r})")

    display_title = title or dest.stem
    mime_type = mimetypes.guess_type(dest.name)[0] or "application/octet-stream"
    ext = dest.suffix.lower()
    attachment_ref = str(dest.relative_to(VAULT_PATH))

    extracted = ""
    status = "attachment only, no text extracted"

    if ext == ".pdf":
        try:
            extracted = extract_pdf_text(dest)
            if extracted:
                status = "PDF text extracted"
                try:
                    tables = extract_pdf_tables(dest)
                    if tables:
                        extracted = f"{extracted}\n\n{tables}"
                        status = "PDF text extracted (+ tables)"
                except Exception as exc:   # table extraction is a bonus on top of already-successful text extraction -- don't lose the text over it
                    logger.warning("PDF table extraction failed for %s: %s", dest, exc)
            else:
                extracted = await extract_pdf_via_ocr(client, dest)
                status = "PDF had no text layer -- OCR'd via vision model" if extracted else "PDF OCR found no text"
        except Exception as exc:   # corrupted/encrypted PDF, or OCR fallback itself failed -- keep the attachment, just skip text extraction
            logger.warning("PDF extraction failed for %s: %s", dest, exc)
            status = f"PDF text extraction failed ({exc}) -- attachment saved anyway"
    elif ext in IMAGE_EXTENSIONS:
        try:
            extracted = await ocr_image(client, dest.read_bytes(), mime_type)
            status = "image OCR'd"
        except Exception as exc:   # vision model AND local OCR fallback both failed -- keep the attachment, just skip extraction
            logger.warning("image OCR failed for %s: %s", dest, exc)
            status = f"image OCR failed ({exc}) -- attachment saved anyway"
    elif ext == ".docx":
        try:
            extracted = extract_docx_text(dest)
            status = "docx text extracted" if extracted else "docx had no extractable text"
        except Exception as exc:   # corrupted/unsupported docx -- keep the attachment, just skip text extraction
            logger.warning("docx extraction failed for %s: %s", dest, exc)
            status = f"docx text extraction failed ({exc}) -- attachment saved anyway"
    else:
        try:
            extracted = dest.read_text(encoding="utf-8")
            status = "read as plain text"
        except (UnicodeDecodeError, OSError):   # binary/unreadable -- fine, attachment-only note
            pass

    link = f"![[{attachment_ref}]]" if ext in IMAGE_EXTENSIONS else f"[[{attachment_ref}]]"
    body = f"{link}\n\n{extracted}" if extracted else link

    path, source = write_note(display_title, tags, audience, body, folder)
    chunk_count = await index_note(client, path, source, display_title, tags, audience)

    return {"status": status, "source": source, "attachment_ref": attachment_ref, "chunk_count": chunk_count}
